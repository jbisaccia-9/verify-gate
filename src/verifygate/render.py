"""Render a Draft to Word and PDF from the *same* object, and read both back.

Both files are built from one Draft so they can't drift -- but the gate does
not trust that. G4 extracts the text of each file independently and compares
them, because a stale PDF sitting next to a fresh .docx is exactly the kind
of thing that gets emailed.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .draft import COMPANY, Draft

TITLES = {
    "verification": "Employment Verification Letter",
    "attestation": "Background Check Attestation",
}


def letter_blocks(d: Draft) -> list[tuple[str, str]]:
    """The letter as an ordered list of (kind, text). Single source for both renderers."""
    blocks = [
        ("company", COMPANY["name"]),
        ("meta", COMPANY["address"]),
        ("meta", f"{COMPANY['hr_email']}  |  {COMPANY['hr_phone']}"),
        ("title", TITLES[d.letter_type]),
        ("meta", d.letter_date),
        ("meta", f"Reference: {d.request_id}"),
        ("body", f"{d.addressee},"),
    ]
    blocks += [("body", p) for p in d.paragraphs]
    blocks += [
        ("body", "Sincerely,"),
        ("sig", COMPANY["signer_name"]),
        ("meta", COMPANY["signer_title"]),
        ("meta", COMPANY["name"]),
    ]
    return blocks


def render_docx(d: Draft, path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for kind, text in letter_blocks(d):
        p = doc.add_paragraph()
        run = p.add_run(text)
        if kind == "company":
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        elif kind == "title":
            run.bold = True
            run.font.size = Pt(13)
            p.paragraph_format.space_before = Pt(14)
        elif kind == "meta":
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif kind == "sig":
            run.bold = True
            p.paragraph_format.space_before = Pt(24)
        if kind == "body":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(10)
    doc.save(path)
    Document(path)  # reopen: a malformed part should fail here, not on the recipient's machine
    return path


def render_pdf(d: Draft, path: Path) -> Path:
    ss = getSampleStyleSheet()
    st = {
        "company": ParagraphStyle("c", parent=ss["Title"], fontSize=16, alignment=0, spaceAfter=2),
        "title": ParagraphStyle("t", parent=ss["Heading2"], fontSize=13, spaceBefore=14),
        "meta": ParagraphStyle("m", parent=ss["Normal"], fontSize=9.5, textColor="#555555"),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=11, leading=15, spaceAfter=10),
        "sig": ParagraphStyle("s", parent=ss["Normal"], fontSize=11, spaceBefore=24,
                              fontName="Helvetica-Bold"),
    }
    story = []
    for kind, text in letter_blocks(d):
        safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe, st[kind]))
        if kind == "meta" and text == f"Reference: {d.request_id}":
            story.append(Spacer(1, 0.2 * inch))
    SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=inch, rightMargin=inch,
                      topMargin=inch, bottomMargin=inch).build(story)
    PdfReader(path)  # reopen
    return path


def render_all(d: Draft, outdir: Path) -> tuple[Path, Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    stem = f"{d.request_id}-{d.letter_type}"
    return render_docx(d, outdir / f"{stem}.docx"), render_pdf(d, outdir / f"{stem}.pdf")


# ------------------------------------------------------------- read-back

def normalize(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u2019", "'")
    text = re.sub(r"-\s*\n\s*", "-", text)         # PDF line-wrap hyphenation
    return re.sub(r"\s+", " ", text).strip()


def text_of_docx(path: Path) -> str:
    return normalize("\n".join(p.text for p in Document(path).paragraphs))


def text_of_pdf(path: Path) -> str:
    return normalize("\n".join(page.extract_text() or "" for page in PdfReader(path).pages))
